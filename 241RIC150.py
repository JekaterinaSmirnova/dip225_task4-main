from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.common.by import By
import time
from openpyxl import load_workbook
from collections import defaultdict
service = Service()
option = webdriver.ChromeOptions()
driver = webdriver.Chrome(service=service, options=option)

name = []

with open("people.csv", "r") as file:
    next(file)  
    for line in file:
        row = line.rstrip().split(",")
        name.append(row[2] + " " + row[3]) 

url = "https://emn178.github.io/online-tools/crc32.html"
driver.get(url)
time.sleep(2)  

crc32_codes = []

for employee_name in name:
    input_box = driver.find_element(By.ID, "input")
    input_box.clear() 
    input_box.send_keys(employee_name)
    output_box = driver.find_element(By.ID, "output")
    crc32_value = output_box.get_attribute("value")
    crc32_codes.append(crc32_value)

wb = load_workbook("salary.xlsx")
ws = wb.active

salary_map = defaultdict(float) 

for row in range(2, ws.max_row + 1): 
    crc32_in_sheet = ws.cell(row=row, column=1).value
    salary = ws.cell(row=row, column=2).value
    if salary is not None: 
        salary = float(salary)  
        salary_map[crc32_in_sheet] += salary  

target_employees = ["Chloe Ramirez", "Hunter Hahn", "Jo Rivers", "Adrienne Lambert"]
for i, crc32_code in enumerate(crc32_codes):
    employee_name = name[i]
    if employee_name in target_employees:  
        if crc32_code in salary_map:
            salary = salary_map[crc32_code]  
            print(f"Darbinieka {employee_name} kopējā alga: {salary}")